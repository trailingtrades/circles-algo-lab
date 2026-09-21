# -*- coding: utf-8 -*-
"""Slide primitives: a 16:9 deck of native, editable PowerPoint shapes (python-pptx), nothing rasterised.

All geometry is in inches on a 13.333 x 7.5 in slide. Text boxes have zero internal margins so text lines up with
the shapes around it, and every box is sized from fit.py's measurements before it is drawn.
"""
from __future__ import annotations

from lxml import etree
from pptx import Presentation
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from . import fit
from .theme import (BG, CARD, DARK_TX2, FONT, FOOTER, INK, INK3, LINE, LOGO, LOGO_WHITE)

W, H = 13.333, 7.5
M = 0.5              # side margin
FOOT_Y = 7.02        # footer band top
CONTENT_BOTTOM = 6.86

LEFT, CENTER, RIGHT = PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.RIGHT
TOP, MIDDLE, BOTTOM = MSO_ANCHOR.TOP, MSO_ANCHOR.MIDDLE, MSO_ANCHOR.BOTTOM


def E(v):
    return Emu(int(round(v * 914400)))


class Canvas:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W)
        self.prs.slide_height = Inches(H)
        self.blank = self.prs.slide_layouts[6]

    # ---------- slides ----------
    def slide(self, bg=BG):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = bg
        return s

    @property
    def count(self):
        return len(self.prs.slides._sldIdLst)

    def notes(self, slide, text: str):
        slide.notes_slide.notes_text_frame.text = text.strip()

    def save(self, path):
        self.prs.save(path)

    # ---------- text ----------
    def text(self, slide, x, y, w, h, paras, size, color=INK, bold=False, italic=False, align=LEFT, anchor=TOP,
             spacing=1.0, gap=0.0, font=FONT, bullet=None, bullet_color=None, indent=0.26, caps=False, name=None):
        """A text box. `paras` is a string (newlines = paragraphs) or a list whose items are strings or lists of
        runs (text, {bold, italic, color, size}). `gap` = points of space after each paragraph. `bullet` = a
        bullet character drawn in `bullet_color`, hanging `indent` inches."""
        tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        items = paras.split("\n") if isinstance(paras, str) else list(paras)
        for i, item in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            if gap and i < len(items) - 1:
                p.space_after = Pt(gap)
            runs = [(item, {})] if isinstance(item, str) else item
            for txt, st in runs:
                r = p.add_run()
                r.text = txt.upper() if caps else txt
                f = r.font
                f.name = font
                f.size = Pt(st.get("size", size))
                f.bold = st.get("bold", bold)
                f.italic = st.get("italic", italic)
                f.color.rgb = st.get("color", color)
            if bullet:
                self._bullet(p, bullet, bullet_color or color, indent)
        return tb

    @staticmethod
    def _bullet(p, char, colour, indent):
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(indent * 914400)))
        pPr.set("indent", str(-int(indent * 914400)))
        for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        clr = etree.SubElement(buClr, qn("a:srgbClr"))
        clr.set("val", str(colour))
        buFont = etree.SubElement(pPr, qn("a:buFont"))
        buFont.set("typeface", "Arial")
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", char)

    # ---------- shapes ----------
    def box(self, slide, x, y, w, h, fill=CARD, line=LINE, lw=0.75, radius=0.1, dash=False, shape=None, name=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
        sh = slide.shapes.add_shape(kind, E(x), E(y), E(w), E(h))
        if name:
            sh.name = name
        if fill is None:
            sh.fill.background()
        else:
            sh.fill.solid()
            sh.fill.fore_color.rgb = fill
        if line is None:
            sh.line.fill.background()
        else:
            sh.line.color.rgb = line
            sh.line.width = Pt(lw)
            if dash:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        sh.shadow.inherit = False
        if kind == MSO_SHAPE.ROUNDED_RECTANGLE and radius:
            try:
                sh.adjustments[0] = max(0.0, min(0.5, radius / max(0.01, min(w, h))))
            except (IndexError, ValueError):
                pass
        if sh.has_text_frame:
            tf = sh.text_frame
            tf.word_wrap = True
            tf.auto_size = MSO_AUTO_SIZE.NONE
        return sh

    def label(self, sh, text, size, color=INK, bold=False, italic=False, align=CENTER, anchor=MIDDLE, margin=0.0,
              spacing=1.0, font=FONT):
        """Text inside a shape (it moves and resizes with the shape when someone edits the deck)."""
        tf = sh.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = E(margin)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        lines = text.split("\n") if isinstance(text, str) else text
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = spacing
            runs = [(ln, {})] if isinstance(ln, str) else ln
            for txt, st in runs:
                r = p.add_run()
                r.text = txt
                r.font.name = font
                r.font.size = Pt(st.get("size", size))
                r.font.bold = st.get("bold", bold)
                r.font.italic = st.get("italic", italic)
                r.font.color.rgb = st.get("color", color)
        return sh

    def pill(self, slide, x, y, text, size, fill, color, bold=True, pad=0.12, h=None, line=None, min_w=0.0):
        """A rounded chip sized to its text. Returns its width."""
        tw = fit.width_in(text, size, bold) / fit.WRAP_SLACK
        w = max(min_w, tw + 2 * pad)
        h = h or (size * 1.2 / 72 + 0.12)
        sh = self.box(slide, x, y, w, h, fill=fill, line=line, radius=h / 2)
        self.label(sh, text, size, color, bold=bold)
        return w

    def oval(self, slide, x, y, d, fill, line=None, lw=1.0):
        return self.box(slide, x, y, d, d, fill=fill, line=line, lw=lw, radius=None, shape=MSO_SHAPE.OVAL)

    def line(self, slide, x1, y1, x2, y2, color=LINE, lw=1.0, dash=False, name=None):
        cx = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        cx.line.color.rgb = color
        cx.line.width = Pt(lw)
        if dash:
            cx.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        # the default connector style carries a theme shadow: switch it off, as for every other shape
        spPr = cx._element.spPr
        if spPr.find(qn("a:effectLst")) is None:
            etree.SubElement(spPr, qn("a:effectLst"))
        if name:
            cx.name = name
        return cx

    def polyline(self, slide, pts, color, lw=2.25, name=None):
        if len(pts) < 2:
            return None
        fb = slide.shapes.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
        fb.add_line_segments([(E(px), E(py)) for px, py in pts[1:]], close=False)
        sh = fb.convert_to_shape()
        sh.fill.background()
        sh.line.color.rgb = color
        sh.line.width = Pt(lw)
        sh.shadow.inherit = False
        # round joins, so a steep zig-zag does not grow spikes
        ln = sh.line._get_or_add_ln()
        for el in ln.findall(qn("a:round")):
            ln.remove(el)
        etree.SubElement(ln, qn("a:round"))
        if name:
            sh.name = name
        return sh

    def logo(self, slide, x, y, size, white=False):
        path = LOGO_WHITE if white else LOGO
        return slide.shapes.add_picture(path, E(x), E(y), height=E(size))

    def footer(self, slide, right: str, dark=False):
        """Logo + credential line on the left, where-am-I + page number on the right. On EVERY slide."""
        col = DARK_TX2 if dark else INK3
        self.logo(slide, M, FOOT_Y + 0.02, 0.26, white=False)
        self.text(slide, M + 0.36, FOOT_Y + 0.07, 6.0, 0.22, FOOTER, 9, col, name="Footer")
        self.text(slide, W - M - 5.0, FOOT_Y + 0.07, 5.0, 0.22, f"{right}   {self.count}", 9, col, align=RIGHT)
