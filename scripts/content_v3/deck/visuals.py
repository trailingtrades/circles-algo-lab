# -*- coding: utf-8 -*-
"""The SMART teaching visuals (apps/learn/src/lib/content/visuals.ts) drawn as native PowerPoint shapes.

One renderer per kind: flow, steps, compare, calc, bars, line, candles, mindmap, table (the story is a slide of its
own, see slides.py). Each draws into a rectangle and adapts to its shape: a wide band gets a left-to-right flow or
timeline, a tall column gets a vertical one. Every number on a candle or line chart is made up, so both carry the
app's "Illustrative" stamp. A table is a real PowerPoint table, so a mentor can type into the template rows.
"""
from __future__ import annotations

import math

from lxml import etree
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.util import Pt

from . import fit
from .canvas import CENTER, E, LEFT, MIDDLE, RIGHT
from .theme import (BEAR, BEAR_VOL, BRAND_TEXT, BULL, BULL_VOL, CARD, CARD2, CYAN, CYAN_SOFT, INK, INK2, INK3,
                    LINE, LINE2, NAVY, TONE_FILL, TONE_SOFT, TONE_TEXT, WHITE, tone, ui)


def L(x, lang):
    """A three-language string in the deck's language (plain strings pass through)."""
    if x is None:
        return ""
    if isinstance(x, str):
        return x
    return x.get(lang) or x.get("en") or ""


def fmt(n):
    """Tick / value formatting, as the app's fmt(): Indian digit grouping from 1,000."""
    if abs(n) >= 1000:
        s = str(int(round(abs(n))))
        head, tail = s[:-3], s[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        return ("-" if n < 0 else "") + ",".join(groups + [tail])
    if float(n).is_integer():
        return str(int(n))
    return f"{n:.2f}" if abs(n) < 10 else f"{n:.1f}"


def num_text(v):
    return str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)


class Frame:
    """Where a visual is drawn: the deck, the slide, the language, and a report label for overflow messages."""

    def __init__(self, cv, slide, lang, where):
        self.cv, self.s, self.lang, self.where = cv, slide, lang, where

    def t(self, x):
        return L(x, self.lang)

    def text(self, *a, **k):
        return self.cv.text(self.s, *a, **k)

    def box(self, *a, **k):
        return self.cv.box(self.s, *a, **k)

    def line(self, *a, **k):
        return self.cv.line(self.s, *a, **k)

    def overflow(self, what, text, w, h):
        fit.need(0, f"{self.where}/{what}", text, w, h)


CHARTS = {"line", "candles"}
PAD = 0.2


def wants_wide(v) -> bool:
    """Visuals that read better as a full-width band than in a side column."""
    k = v.get("kind")
    if k == "flow":
        return len(v.get("nodes") or []) >= 5
    if k == "steps":
        return len(v.get("items") or []) >= 5
    if k == "compare":
        return len(v.get("cols") or []) >= 3
    if k == "candles":
        return not v.get("anatomy") and len(v.get("bars") or []) > 30
    if k == "table":
        return len(v.get("head") or []) >= 4
    return False


def frame(v, lang, w, h, show_title=True, card=True):
    """The geometry render() will use, without drawing: title and caption sizes and the body rectangle.

    Returned offsets are relative to the frame's top-left corner, so slides can ask "will this fit?" first."""
    iw = w - 2 * PAD
    top = 0.16 if card else 0.0
    bottom = h - (0.14 if card else 0.0)
    g = {"iw": iw, "title": "", "tpt": 0, "th": 0, "illus_w": 0.0, "cap": "", "cpt": 0, "ch": 0}
    title = L(v.get("title"), lang) if show_title else ""
    chart = v.get("kind") in CHARTS
    g["title_y"] = top
    if title or chart:
        if chart:
            g["illus_w"] = fit.width_in(ui("illustrative", lang), 9, True) + 0.3
        if title:
            tw = iw - g["illus_w"] - 0.15
            g["tpt"] = fit.fit(title, tw, 0.5, 13.5, 10.5, bold=True) or 10.5
            g["th"] = fit.height(title, g["tpt"], tw, bold=True)
            g["title"], g["tw"] = title, tw
            top += g["th"] + 0.12
        else:
            top += 0.34
    cap = L(v.get("caption"), lang)
    if cap:
        cap_h_max = 0.52 if h < 3.2 else 0.8
        g["cpt"] = fit.fit(cap, iw, cap_h_max, 10.5, 9) or 9
        g["cap"] = fit.elide(cap, iw, cap_h_max, g["cpt"])
        g["ch"] = fit.height(g["cap"], g["cpt"], iw)
        g["cap_y"] = bottom - g["ch"]
        bottom -= g["ch"] + 0.1
    g["body_y"], g["body_h"] = top, max(0.4, bottom - top)
    return g


def body_size(v, lang, w, h, show_title=True, card=True):
    """(width, height) of the area a renderer gets inside a w x h frame."""
    g = frame(v, lang, w, h, show_title, card)
    return g["iw"], g["body_h"]


def render(cv, slide, v, x, y, w, h, lang, where, show_title=True, card=True):
    """Draw visual v in the rectangle, on a white card. Returns nothing; overflows are recorded in fit.OVERFLOW."""
    f = Frame(cv, slide, lang, where)
    if card:
        f.box(x, y, w, h, fill=CARD, line=LINE, radius=0.12)
    g = frame(v, lang, w, h, show_title, card)
    ix, iw = x + PAD, g["iw"]
    if g["illus_w"]:
        cv.pill(slide, ix + iw - g["illus_w"], y + g["title_y"] + 0.01, ui("illustrative", lang), 9,
                TONE_SOFT["gold"], TONE_TEXT["gold"], pad=0.12)
    if g["title"]:
        f.text(ix, y + g["title_y"], g["tw"], g["th"] + 0.02, g["title"], g["tpt"], INK, bold=True)
    if g["cap"]:
        f.text(ix, y + g["cap_y"], iw, g["ch"] + 0.02, g["cap"], g["cpt"], INK2)
    fn = RENDERERS.get(v.get("kind"))
    if not fn:
        fit.OVERFLOW.append(f"{where}: no renderer for visual kind '{v.get('kind')}'")
        return
    fn(f, v, ix, y + g["body_y"], iw, g["body_h"])


# ------------------------------------------------------------------ calc
def v_calc(f, v, x, y, w, h):
    rows = list(v.get("rows") or [])
    res = v.get("result") or {}
    n = len(rows) + 1
    labels = [f.t(r.get("label")) for r in rows] + [f.t(res.get("label"))]
    values = [str(r.get("value", "")) for r in rows] + [str(res.get("value", ""))]
    ops = [r.get("op") or "" for r in rows] + ["="]
    opw = 0.36
    rh = min(0.66, (h - 0.08) / n)
    for pt in [x / 2 for x in range(30, 19, -1)]:  # 15 .. 10
        valw = min(w * 0.46, max(fit.width_in(s, pt + 1, True) for s in values) / fit.WRAP_SLACK + 0.12)
        lw = w - opw - valw - 0.12
        if all(fit.height(lab, pt, lw, bold=(i == n - 1)) <= rh - 0.1 for i, lab in enumerate(labels)) and \
                all(fit.fits_one_line(s, pt + 1, valw, True) for s in values):
            break
    else:
        f.overflow("calc", " | ".join(labels), w, h)
    total = n * rh + 0.06
    y0 = y + max(0.0, (h - total) / 2)
    for i in range(n):
        last = i == n - 1
        ry = y0 + i * rh + (0.06 if last else 0)
        if last:
            f.box(x, ry, w, rh, fill=TONE_SOFT["brand"], line=None, radius=0.08)
        f.text(x, ry, opw, rh, ops[i], pt + 1, INK3, bold=True, align=CENTER, anchor=MIDDLE)
        f.text(x + opw, ry, lw, rh, labels[i], pt, INK, bold=last, anchor=MIDDLE)
        f.text(x + w - valw - 0.08, ry, valw, rh, values[i], pt + (1.5 if last else 1), BRAND_TEXT if last else INK,
               bold=True, align=RIGHT, anchor=MIDDLE)
        if not last and i < n - 2:
            f.line(x, ry + rh, x + w, ry + rh, LINE, 0.75)


# ------------------------------------------------------------------ bars
def v_bars(f, v, x, y, w, h):
    items = list(v.get("items") or [])
    n = max(1, len(items))
    unit = v.get("unit") or ""
    labels = [f.t(b.get("label")) for b in items]
    vals = [b.get("display") if b.get("display") is not None else f"{num_text(b.get('value', 0))}{unit}" for b in items]
    mx = max([abs(float(b.get("value", 0))) for b in items] + [1e-9])
    rh = min(0.72, h / n)
    for pt in [x / 2 for x in range(28, 19, -1)]:  # 14 .. 10
        lw = min(w * 0.42, max(fit.width_in(s, pt) for s in labels) / fit.WRAP_SLACK + 0.08)
        lw = max(lw, min(w * 0.42, 1.2))
        vw = max(fit.width_in(s, pt, True) for s in vals) / fit.WRAP_SLACK + 0.1
        if all(fit.height(s, pt, lw) <= rh - 0.06 for s in labels):
            break
    else:
        f.overflow("bars", " | ".join(labels), w, h)
    tw = w - lw - vw - 0.24
    bt = min(0.34, rh * 0.5)
    y0 = y + max(0.0, (h - n * rh) / 2)
    for i, b in enumerate(items):
        ry = y0 + i * rh
        f.text(x, ry, lw, rh, labels[i], pt, INK, anchor=MIDDLE)
        tx = x + lw + 0.12
        f.box(tx, ry + (rh - bt) / 2, tw, bt, fill=CARD2, line=None, radius=bt / 2)
        fw = max(0.05, tw * abs(float(b.get("value", 0))) / mx)
        f.box(tx, ry + (rh - bt) / 2, fw, bt, fill=TONE_FILL[tone(b.get("tone"))], line=None, radius=bt / 2)
        f.text(tx + tw + 0.1, ry, vw + 0.05, rh, vals[i], pt, TONE_TEXT[tone(b.get("tone"))], bold=True, anchor=MIDDLE)


# ------------------------------------------------------------------ compare
def v_compare(f, v, x, y, w, h):
    cols = list(v.get("cols") or [])
    n = max(1, len(cols))
    gap = 0.18
    cw = (w - gap * (n - 1)) / n
    inner = cw - 0.32
    heads = [f.t(c.get("head")) for c in cols]
    pts = [[f.t(p) for p in (c.get("points") or [])] for c in cols]
    chosen = None
    for pt in [x / 2 for x in range(30, 19, -1)]:  # 15 .. 10
        hpt = pt + 1
        head_h = max(fit.height(s, hpt, inner, bold=True) for s in heads) + 0.24
        body_h = max(fit.paras_height(p, pt, inner - 0.22, 1.0, 0.5 * pt) for p in pts) + 0.3
        if head_h + body_h <= h:
            chosen = (pt, hpt, head_h, body_h)
            break
    if not chosen:
        pt, hpt = 10, 11
        head_h = max(fit.height(s, hpt, inner, bold=True) for s in heads) + 0.24
        body_h = h - head_h
        f.overflow("compare", " | ".join(heads), w, h)
    else:
        pt, hpt, head_h, body_h = chosen
    ch = min(h, head_h + body_h + 0.25)
    y0 = y + (h - ch) / 2
    for i, c in enumerate(cols):
        t = tone(c.get("tone"))
        cx = x + i * (cw + gap)
        f.box(cx, y0, cw, ch, fill=CARD, line=None, radius=0.1)
        f.box(cx, y0, cw, head_h, fill=TONE_SOFT[t], line=None, radius=0.1)
        # square off the header's lower corners so it sits flush on the card body, then outline the whole card
        f.box(cx, y0 + head_h / 2, cw, head_h / 2, fill=TONE_SOFT[t], line=None, radius=None)
        f.box(cx, y0, cw, ch, fill=None, line=TONE_FILL[t], lw=1.25, radius=0.1)
        f.text(cx + 0.16, y0, inner, head_h, heads[i], hpt, TONE_TEXT[t], bold=True, anchor=MIDDLE)
        f.line(cx, y0 + head_h, cx + cw, y0 + head_h, TONE_FILL[t], 1.25)
        f.text(cx + 0.16, y0 + head_h + 0.16, inner, ch - head_h - 0.2, pts[i], pt, INK, gap=0.5 * pt,
               bullet="•", bullet_color=TONE_FILL[t], indent=0.22)


# ------------------------------------------------------------------ steps
def v_steps(f, v, x, y, w, h):
    items = list(v.get("items") or [])
    if w / max(h, 0.1) > 2.4 and len(items) >= 3:
        return _steps_row(f, items, x, y, w, h)
    n = max(1, len(items))
    tags = [s.get("tag") or str(i + 1) for i, s in enumerate(items)]
    dw = max(0.46, max(fit.width_in(t, 11, True) for t in tags) + 0.24)
    dh = 0.36
    tx = x + dw + 0.2
    tw = w - dw - 0.2
    labs = [f.t(s.get("label")) for s in items]
    subs = [f.t(s.get("sub")) for s in items]
    gap = 0.14
    chosen = None
    for pt in [x / 2 for x in range(28, 19, -1)]:  # 14 .. 10
        spt = pt - 1.5
        hs = [max(dh, fit.height(labs[i], pt, tw, bold=True) + (fit.height(subs[i], spt, tw) if subs[i] else 0))
              for i in range(n)]
        if sum(hs) + gap * (n - 1) <= h:
            chosen = (pt, spt, hs)
            break
    if not chosen:
        pt, spt = 10, 8.5
        hs = [max(dh, fit.height(labs[i], pt, tw, bold=True) + fit.height(subs[i], spt, tw)) for i in range(n)]
        f.overflow("steps", " | ".join(labs), w, h)
    else:
        pt, spt, hs = chosen
    total = sum(hs) + gap * (n - 1)
    extra = max(0.0, h - total)
    gap2 = gap + min(0.25, extra / max(1, n - 1)) if n > 1 else gap
    total = sum(hs) + gap2 * (n - 1)
    y0 = y + max(0.0, (h - total) / 2)
    ys, cy = [], y0
    for hh in hs:
        ys.append(cy)
        cy += hh + gap2
    if n > 1:
        f.line(x + dw / 2, ys[0] + dh / 2, x + dw / 2, ys[-1] + dh / 2, LINE2, 1.5)
    for i in range(n):
        d = f.box(x, ys[i], dw, dh, fill=CYAN_SOFT, line=CYAN, lw=1.25, radius=dh / 2)
        f.cv.label(d, tags[i], 11, BRAND_TEXT, bold=True)
        paras = [[(labs[i], {"bold": True, "size": pt, "color": INK})]]
        if subs[i]:
            paras.append([(subs[i], {"size": spt, "color": INK2})])
        f.text(tx, ys[i] + (0.02 if hs[i] > dh else (dh - fit.line_h(pt)) / 2), tw, hs[i], paras, pt)


def _steps_row(f, items, x, y, w, h):
    n = len(items)
    cw = w / n
    inner = cw - 0.16
    tags = [s.get("tag") or str(i + 1) for i, s in enumerate(items)]
    dw = max(0.46, max(fit.width_in(t, 11, True) for t in tags) + 0.24)
    dh = 0.36
    labs = [f.t(s.get("label")) for s in items]
    subs = [f.t(s.get("sub")) for s in items]
    avail = h - dh - 0.14
    chosen = None
    for pt in [x / 2 for x in range(28, 18, -1)]:  # 14 .. 9.5
        spt = pt - 1.5
        need = max(fit.height(labs[i], pt, inner, bold=True) + (fit.height(subs[i], spt, inner) + 0.04 if subs[i] else 0)
                   for i in range(n))
        if need <= avail:
            chosen = (pt, spt, need)
            break
    if not chosen:
        pt, spt, need = 9.5, 8.5, avail
        f.overflow("steps-row", " | ".join(labs), w, h)
    else:
        pt, spt, need = chosen
    total = dh + 0.14 + need
    y0 = y + max(0.0, (h - total) / 2)
    f.line(x + cw / 2, y0 + dh / 2, x + w - cw / 2, y0 + dh / 2, LINE2, 1.5)
    for i in range(n):
        cx = x + i * cw
        d = f.box(cx + (cw - dw) / 2, y0, dw, dh, fill=CYAN_SOFT, line=CYAN, lw=1.25, radius=dh / 2)
        f.cv.label(d, tags[i], 11, BRAND_TEXT, bold=True)
        paras = [[(labs[i], {"bold": True, "size": pt, "color": INK})]]
        if subs[i]:
            paras.append([(subs[i], {"size": spt, "color": INK2})])
        f.text(cx + 0.08, y0 + dh + 0.14, inner, need + 0.05, paras, pt, align=CENTER, gap=2)


# ------------------------------------------------------------------ flow
def v_flow(f, v, x, y, w, h):
    nodes = list(v.get("nodes") or [])
    arrows = [f.t(a) for a in (v.get("arrows") or [])]
    if w / max(h, 0.1) > 2.2:
        return _flow_row(f, nodes, arrows, x, y, w, h)
    n = max(1, len(nodes))
    labs = [f.t(nd.get("label")) for nd in nodes]
    subs = [f.t(nd.get("sub")) for nd in nodes]
    inner = w - 0.3
    ag = 0.36
    chosen = None
    for pt in [x / 2 for x in range(28, 19, -1)]:  # 14 .. 10
        spt = pt - 1.5
        hs = [fit.height(labs[i], pt, inner, bold=True) + (fit.height(subs[i], spt, inner) if subs[i] else 0) + 0.18
              for i in range(n)]
        if sum(hs) + ag * (n - 1) <= h:
            chosen = (pt, spt, hs)
            break
    if not chosen:
        pt, spt = 10, 8.5
        hs = [fit.height(labs[i], pt, inner, bold=True) + fit.height(subs[i], spt, inner) + 0.16 for i in range(n)]
        f.overflow("flow", " | ".join(labs), w, h)
    else:
        pt, spt, hs = chosen
    extra = max(0.0, h - sum(hs) - ag * (n - 1))
    ag2 = ag + min(0.2, extra / max(1, n - 1)) if n > 1 else ag
    y0 = y + max(0.0, (h - sum(hs) - ag2 * (n - 1)) / 2)
    cy = y0
    for i in range(n):
        t = tone(nodes[i].get("tone"))
        f.box(x, cy, w, hs[i], fill=CARD, line=TONE_FILL[t], lw=1.5, radius=0.1)
        paras = [[(labs[i], {"bold": True, "size": pt, "color": INK})]]
        if subs[i]:
            paras.append([(subs[i], {"size": spt, "color": INK2})])
        f.text(x + 0.15, cy + 0.09, inner, hs[i] - 0.1, paras, pt)
        cy += hs[i]
        if i < n - 1:
            ah = ag2 - 0.08
            a = f.box(x + 0.34, cy + 0.04, 0.24, ah, fill=LINE2, line=None, radius=None, shape=MSO_SHAPE.DOWN_ARROW)
            if i < len(arrows) and arrows[i]:
                apt = fit.fit_single_line(arrows[i], w - 0.8, 10.5, 8.5) or 8.5
                f.text(x + 0.7, cy, w - 0.8, ag2, arrows[i], apt, INK3, italic=True, anchor=MIDDLE)
            cy += ag2


def _flow_row(f, nodes, arrows, x, y, w, h):
    n = max(1, len(nodes))
    has_lab = any(arrows)
    ag = 0.78 if has_lab else 0.42
    nw = (w - ag * (n - 1)) / n
    inner = nw - 0.24
    labs = [f.t(nd.get("label")) for nd in nodes]
    subs = [f.t(nd.get("sub")) for nd in nodes]
    chosen = None
    for pt in [x / 2 for x in range(27, 17, -1)]:  # 13.5 .. 9
        spt = pt - 1.5
        need = max(fit.height(labs[i], pt, inner, bold=True) + (fit.height(subs[i], spt, inner) if subs[i] else 0)
                   for i in range(n)) + 0.24
        if need <= h and all(fit.width_in(wd, pt, True) <= inner for s in labs for wd in s.split()):
            chosen = (pt, spt, need)
            break
    if not chosen:
        pt, spt, need = 9, 8, h
        f.overflow("flow-row", " | ".join(labs), w, h)
    else:
        pt, spt, need = chosen
    nh = min(h, max(need, 0.9))
    y0 = y + (h - nh) / 2
    for i in range(n):
        t = tone(nodes[i].get("tone"))
        nx = x + i * (nw + ag)
        f.box(nx, y0, nw, nh, fill=CARD, line=TONE_FILL[t], lw=1.5, radius=0.1)
        paras = [[(labs[i], {"bold": True, "size": pt, "color": INK})]]
        if subs[i]:
            paras.append([(subs[i], {"size": spt, "color": INK2})])
        f.text(nx + 0.12, y0, inner, nh, paras, pt, align=CENTER, anchor=MIDDLE)
        if i < n - 1:
            ax = nx + nw + 0.08
            f.box(ax + (ag - 0.16 - 0.3) / 2, y0 + nh / 2 - 0.1, 0.3, 0.2, fill=LINE2, line=None, radius=None,
                  shape=MSO_SHAPE.RIGHT_ARROW)
            if i < len(arrows) and arrows[i]:
                lab = arrows[i]
                aw = ag - 0.06
                apt = next((p for p in (9, 8.5, 8) if fit.height(lab, p, aw, italic=True) <= nh / 2 - 0.12), 8)
                lh = fit.height(lab, apt, aw, italic=True)
                f.text(nx + nw + 0.03, y0 + nh / 2 - 0.14 - lh, aw, lh + 0.02, lab, apt, INK3, italic=True,
                       align=CENTER)


# ------------------------------------------------------------------ line
def v_line(f, v, x, y, w, h):
    series = [s for s in (v.get("series") or []) if s.get("points")]
    if not series:
        return
    defaults = ["brand", "down", "gold"]
    cols = [TONE_FILL[tone(s.get("tone"), defaults[i % 3])] for i, s in enumerate(series)]
    allp = [p for s in series for p in s["points"]]
    lo, hi = min(allp), max(allp)
    pad = (hi - lo) * 0.08 or abs(hi) * 0.1 or 1
    mn, mxv = lo - pad, hi + pad
    npts = max(len(s["points"]) for s in series)
    tpt = 9.5
    ticks = [mn + (mxv - mn) * i / 4 for i in range(5)]
    tick_s = [fmt(t) for t in ticks]
    gut = max(fit.width_in(s, tpt) for s in tick_s) + 0.14
    top = y
    yl, xl = f.t(v.get("y_label")), f.t(v.get("x_label"))
    if len(series) > 1:  # legend
        lx = x + gut
        for i, s in enumerate(series):
            nm = f.t(s.get("name"))
            f.box(lx, top + 0.07, 0.28, 0.07, fill=cols[i], line=None, radius=None)
            nw = fit.width_in(nm, 10) / fit.WRAP_SLACK + 0.05
            f.text(lx + 0.36, top, nw, 0.22, nm, 10, INK2)
            lx += 0.36 + nw + 0.3
        top += 0.3
    if yl:
        f.text(x, top, gut + 2.5, 0.2, yl, 9, INK3, bold=True)
        top += 0.24
    px0 = x + gut
    pw = w - gut - 0.08
    marks = [m for m in (v.get("marks") or []) if isinstance(m.get("x"), (int, float))]
    xpos = lambda i: px0 + (0 if npts <= 1 else i / (npts - 1) * pw)
    mpt = 10
    rows, placed = [], []
    for m in marks:
        txt = f.t(m.get("label"))
        mw = fit.width_in(txt, mpt, True) / fit.WRAP_SLACK + 0.04
        mx = xpos(m["x"])
        tx, right = mx + 0.05, False
        if tx + mw > px0 + pw:
            tx, right = mx - 0.05 - mw, True
        r = 0
        while any(rr == r and tx < b + 0.08 and tx + mw > a - 0.08 for a, b, rr in rows):
            r += 1
        rows.append((tx, tx + mw, r))
        placed.append((mx, tx, mw, r, txt))
    nrows = (max(r for *_, r in rows) + 1) if rows else 0
    rowh = 0.22
    mark_top = top
    top += nrows * rowh + (0.04 if nrows else 0)
    bottom = y + h - (0.26 if xl else 0.06)
    ph = max(0.4, bottom - top)
    fy = lambda p: top + (mxv - p) / (mxv - mn) * ph
    for i, t in enumerate(ticks):
        gy = fy(t)
        f.line(px0, gy, px0 + pw, gy, LINE, 0.5)
        f.text(x, gy - 0.09, gut - 0.1, 0.18, tick_s[i], tpt, INK3, align=RIGHT)
    for mx, tx, mw, r, txt in placed:
        ty = mark_top + r * rowh
        f.line(mx, ty + 0.1, mx, top + ph, INK3, 0.75, dash=True)
        f.text(tx, ty, mw, rowh, txt, mpt, INK, bold=True)
    for i, s in enumerate(series):
        f.cv.polyline(f.s, [(xpos(j), fy(p)) for j, p in enumerate(s["points"])], cols[i], 2.25)
    if xl:
        f.text(px0, y + h - 0.22, pw, 0.2, xl, 9, INK3, bold=True, align=CENTER)


# ------------------------------------------------------------------ candles
def v_candles(f, v, x, y, w, h):
    bars = [b for b in (v.get("bars") or []) if isinstance(b, list) and len(b) == 4]
    if not bars:
        return
    if v.get("anatomy"):
        return _anatomy(f, v, bars, x, y, w, h)
    levels = list(v.get("levels") or [])
    marks = [m for m in (v.get("marks") or []) if isinstance(m.get("i"), int) and 0 <= m["i"] < len(bars)]
    vol = v.get("volume") if v.get("volume") and len(v["volume"]) == len(bars) else None
    lo = min([b[2] for b in bars] + [lv["y"] for lv in levels])
    hi = max([b[1] for b in bars] + [lv["y"] for lv in levels])
    pad = (hi - lo) * 0.08 or 1
    tpt, lpt, mpt = 9.5, 10, 10
    ticks = [lo - pad + (hi - lo + 2 * pad) * i / 3 for i in range(4)]
    tick_s = [fmt(t) for t in ticks]
    vol_lab = ui("volume", f.lang) if vol else ""
    gut = max([fit.width_in(s, tpt) for s in tick_s] + [fit.width_in(vol_lab, tpt)]) + 0.14
    lv_txt = [f.t(lv.get("label")) for lv in levels]
    rg = min(max([fit.width_in(s, lpt, True) / fit.WRAP_SLACK for s in lv_txt] + [0]) + 0.16, w * 0.38) if levels else 0.08
    px0, pw = x + gut, w - gut - rg
    step = pw / len(bars)
    cw = max(min(0.05, step * 0.8), min(0.3, step * 0.6))
    cx = lambda i: px0 + step * (i + 0.5)
    # mark labels: one row per clash, above the high (default) or below the low
    placed, rows = [], {"high": [], "low": []}
    for m in marks:
        side = "low" if m.get("at") == "low" else "high"
        txt = f.t(m.get("label"))
        mw = fit.width_in(txt, mpt, True) / fit.WRAP_SLACK + 0.04
        tx = min(max(cx(m["i"]) - mw / 2, px0 - gut + 0.02), px0 + pw + rg - mw)
        r = 0
        while any(rr == r and tx < b + 0.06 and tx + mw > a - 0.06 for a, b, rr in rows[side]):
            r += 1
        rows[side].append((tx, tx + mw, r))
        placed.append((m, side, txt, tx, mw, r))
    rowh = 0.21
    n_hi = max([r for *_, r in rows["high"]] + [-1]) + 1
    n_lo = max([r for *_, r in rows["low"]] + [-1]) + 1
    top_pad = 0.1 + (0.12 + n_hi * rowh if n_hi else 0)
    bot_pad = 0.1 + (0.12 + n_lo * rowh if n_lo else 0)
    vh = min(0.6, h * 0.2) if vol else 0
    py0 = y + top_pad
    ph = max(0.5, h - top_pad - bot_pad - vh)
    pmin, pmax = lo - pad, hi + pad
    fy = lambda p: py0 + (pmax - p) / (pmax - pmin) * ph
    vol_y = py0 + ph + bot_pad   # volume sits under the low-mark labels, never behind them
    for i, t in enumerate(ticks):
        gy = fy(t)
        f.line(px0, gy, px0 + pw, gy, LINE, 0.5)
        f.text(x, gy - 0.09, gut - 0.1, 0.18, tick_s[i], tpt, INK3, align=RIGHT)
    # levels, labels pushed apart in the right gutter
    lv = sorted(({"y": fy(l["y"]), "l": l, "txt": lv_txt[i]} for i, l in enumerate(levels)), key=lambda d: d["y"])
    prev = -1e9
    for d in lv:
        d["ty"] = max(d["y"], prev + 0.2)
        prev = d["ty"]
    for d in lv:
        t = tone(d["l"].get("tone"), "gold")
        f.line(px0, d["y"], px0 + pw + 0.04, d["y"], TONE_FILL[t], 1.5, dash=d["l"].get("dashed") is not False)
        f.text(px0 + pw + 0.08, d["ty"] - 0.1, rg - 0.08, 0.2, d["txt"], lpt, TONE_TEXT[t], bold=True, anchor=MIDDLE)
    vmax = max(vol or [1]) or 1
    for i, (o, hh, ll, c) in enumerate(bars):
        up = c >= o
        col = BULL if up else BEAR
        f.line(cx(i), fy(hh), cx(i), fy(ll), col, 1.25)
        top_ = fy(max(o, c))
        bh = max(0.02, fy(min(o, c)) - top_)
        f.box(cx(i) - cw / 2, top_, cw, bh, fill=col, line=None, radius=None)
        if vol:
            vb = max(0.01, vol[i] / vmax * (vh - 0.04))
            f.box(cx(i) - cw / 2, vol_y + vh - vb, cw, vb, fill=BULL_VOL if up else BEAR_VOL, line=None,
                  radius=None)
    if vol:
        f.text(x, vol_y + vh / 2 - 0.09, gut - 0.1, 0.18, vol_lab, tpt, INK3, align=RIGHT)
    for m, side, txt, tx, mw, r in placed:
        b = bars[m["i"]]
        t = tone(m.get("tone"))
        tri = 0.1
        if side == "high":
            ty = fy(b[1]) - 0.05 - tri
            s = f.box(cx(m["i"]) - 0.06, ty, 0.12, tri, fill=TONE_FILL[t], line=None, radius=None,
                      shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            s.rotation = 180
            ly = ty - 0.02 - rowh * (r + 1)
        else:
            ty = fy(b[2]) + 0.05
            f.box(cx(m["i"]) - 0.06, ty, 0.12, tri, fill=TONE_FILL[t], line=None, radius=None,
                  shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            ly = ty + tri + 0.02 + rowh * r
        f.text(tx, ly, mw, rowh, txt, mpt, TONE_TEXT[t], bold=True, align=CENTER, anchor=MIDDLE)


def _anatomy(f, v, bars, x, y, w, h):
    lang = f.lang
    note = ui("bullish" if bars[0][3] >= bars[0][0] else "bearish", lang)
    nh = 0.3
    ph = h - nh - 0.1
    lo = min(b[2] for b in bars)
    hi = max(b[1] for b in bars)
    pad = (hi - lo) * 0.1 or 1
    fy = lambda p: y + 0.1 + (hi + pad - p) / (hi - lo + 2 * pad) * (ph - 0.2)
    cw = 0.55
    n = len(bars)
    xs = [x + w / 2 + (i - (n - 1) / 2) * 1.1 for i in range(n)]
    for i, (o, hh, ll, c) in enumerate(bars):
        col = BULL if c >= o else BEAR
        f.box(xs[i] - 0.03, fy(hh), 0.06, fy(ll) - fy(hh), fill=col, line=None, radius=None)
        f.box(xs[i] - cw / 2, fy(max(o, c)), cw, max(0.03, fy(min(o, c)) - fy(max(o, c))), fill=col, line=None,
              radius=0.04)
    o, hh, ll, c = bars[0]
    ax = xs[0]
    up = c >= o
    right = [(hh, f"{ui('high', lang)} {fmt(hh)}"), (max(o, c), f"{ui('close' if up else 'open', lang)} {fmt(max(o, c))}"),
             (min(o, c), f"{ui('open' if up else 'close', lang)} {fmt(min(o, c))}"), (ll, f"{ui('low', lang)} {fmt(ll)}")]
    left = [((hh + max(o, c)) / 2, ui("upper_wick", lang)), ((o + c) / 2, ui("body", lang)),
            ((ll + min(o, c)) / 2, ui("lower_wick", lang))]
    apt = 12
    lead = min(1.0, w / 2 - cw / 2 - 0.1 - max(fit.width_in(t, apt, True) for _, t in right + left) - 0.1)
    lead = max(0.25, lead)

    def spread(items):
        out, prev = [], -1e9
        for p, t in sorted(items, key=lambda it: -it[0]):
            yy = max(fy(p), prev + 0.26)
            out.append((fy(p), yy, t))
            prev = yy
        return out

    for py, ty, t in spread(right):
        x1 = ax + cw / 2 + 0.06
        tw = fit.width_in(t, apt, True) / fit.WRAP_SLACK + 0.05
        f.line(x1, py, x1 + lead, ty, INK3, 0.75)
        f.text(x1 + lead + 0.06, ty - 0.12, tw, 0.24, t, apt, INK, bold=True, anchor=MIDDLE)
    for py, ty, t in spread(left):
        x1 = ax - cw / 2 - 0.06
        f.line(x1, py, x1 - lead, ty, INK3, 0.75)
        tw = fit.width_in(t, apt, True) / fit.WRAP_SLACK + 0.05
        f.text(x1 - lead - 0.06 - tw, ty - 0.12, tw, 0.24, t, apt, INK2, bold=True, align=RIGHT, anchor=MIDDLE)
    f.text(x, y + h - nh, w, nh, note, 11, INK2, italic=True, align=CENTER, anchor=MIDDLE)


# ------------------------------------------------------------------ table
def _set_borders(cell, colour, w_pt=0.75):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(int(w_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr"), val=str(colour))
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tcPr.insert(i, ln)


def _cell_text(cell, text, pt, color, bold=False, align=LEFT, italic=False):
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    r = p.add_run()
    r.text = text
    r.font.name = "Segoe UI"
    r.font.size = Pt(pt)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color


MX, MY = 0.07, 0.035   # table cell margins


def _table_cells(v, lang):
    head = [L(c, lang) for c in (v.get("head") or [])]
    nc = len(head)
    rows, plain = [], []
    for r in (v.get("rows") or []):
        r = list(r or [])[:nc] + [""] * max(0, nc - len(r or []))
        rows.append(["" if c is None else (c if isinstance(c, str) else L(c, lang)) for c in r])
        plain.append([isinstance(c, str) for c in r])
    return head, rows, plain


def _table_layout(head, rows, pt, nblank, w):
    """Column widths and row heights for one size, or None when the columns cannot fit the width."""
    nc = len(head)
    nat = []
    for j in range(nc):
        wmax = max([fit.width_in(head[j], pt, True)] + [fit.width_in(r[j], pt) for r in rows])
        nat.append(max(0.8, wmax / fit.WRAP_SLACK + 2 * MX + 0.06))
    if sum(nat) <= w:
        cws = [n * w / sum(nat) for n in nat]
    else:  # squeeze the wide columns first, never below what their longest word needs
        floor = [max(0.8, max(fit.width_in(wd, pt, True) for s in [head[j]] + [r[j] for r in rows]
                               for wd in (s.split() or [""])) / fit.WRAP_SLACK + 2 * MX + 0.05) for j in range(nc)]
        cws = list(nat)
        for _ in range(80):
            over = sum(cws) - w
            if over <= 0.001:
                break
            flex = [j for j in range(nc) if cws[j] > floor[j] + 0.01]
            if not flex:
                break
            share = over / len(flex)
            for j in flex:
                cws[j] = max(floor[j], cws[j] - share)
        if sum(cws) > w + 0.01:
            return None
    hh = max(fit.height(head[j], pt, cws[j] - 2 * MX, bold=True) for j in range(nc)) + 2 * MY + 0.06
    rhs = [max(max(fit.height(r[j], pt, cws[j] - 2 * MX) for j in range(nc)), fit.line_h(pt)) + 2 * MY + 0.05
           for r in rows]
    bh = max(0.36, fit.line_h(pt) + 0.2)
    return {"pt": pt, "nblank": nblank, "cws": cws, "hh": hh, "rhs": rhs, "bh": bh,
            "total": hh + sum(rhs) + nblank * bh}


def table_plan(v, lang, w, h, pt_max=15, pt_min=9, keep_blank=False):
    """The largest size (and as many blank rows as fit) at which the table and its note fit w x h; None if none.
    keep_blank: only a plan with every blank row counts (a template's writing rows are its point)."""
    head, rows, _ = _table_cells(v, lang)
    note = L(v.get("note"), lang)
    npt, note_h = 9, 0.0
    if note:
        npt = fit.fit(note, w, 0.5, 10.5, 9, italic=True) or 9
        note_h = fit.height(note, npt, w, italic=True) + 0.1
    blank = max(0, min(12, int(v.get("blank_rows") or 0)))
    for nblank in range(blank, (blank - 1) if keep_blank else -1, -1):
        for pt in [x / 2 for x in range(int(pt_max * 2), int(pt_min * 2) - 1, -1)]:
            lay = _table_layout(head, rows, pt, nblank, w)
            if lay and lay["total"] <= h - note_h:
                lay.update(note=note, npt=npt, note_h=note_h)
                return lay
    return None


def v_table(f, v, x, y, w, h):
    head, rows, plain = _table_cells(v, f.lang)
    nc = len(head)
    lay = table_plan(v, f.lang, w, h)
    if not lay:
        f.overflow("table", " | ".join(head), w, h)
        lay = _table_layout(head, rows, 9, 0, w) or {"pt": 9, "nblank": 0, "cws": [w / nc] * nc, "hh": 0.4,
                                                     "rhs": [0.3] * len(rows), "bh": 0.36, "total": h}
        lay.update(note=L(v.get("note"), f.lang), npt=9, note_h=0.3)
    pt, nblank, cws, hh, rhs, bh = lay["pt"], lay["nblank"], lay["cws"], lay["hh"], list(lay["rhs"]), lay["bh"]
    avail = h - lay["note_h"]
    # spare height becomes writing room: taller blank rows first, then a little air in the filled rows
    spare = avail - lay["total"]
    if spare > 0 and nblank and bh < 0.62:
        grow = min(spare / nblank, 0.62 - bh)
        bh += grow
        spare -= grow * nblank
    if spare > 0 and rows:
        grow = min(spare / (len(rows) + 1), 0.16)
        rhs = [r + grow for r in rhs]
        hh += grow
    total = hh + sum(rhs) + nblank * bh
    nr = 1 + len(rows) + nblank
    gf = f.s.shapes.add_table(nr, nc, E(x), E(y), E(sum(cws)), E(total))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for j, cwid in enumerate(cws):
        tbl.columns[j].width = E(cwid)
    for i, rh in enumerate([hh] + rhs + [bh] * nblank):
        tbl.rows[i].height = E(rh)
    for i in range(nr):
        for j in range(nc):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = E(MX)
            cell.margin_top = cell.margin_bottom = E(MY)
            cell.vertical_anchor = MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = NAVY
                _cell_text(cell, head[j], pt, WHITE, bold=True)
            elif i <= len(rows):
                cell.fill.fore_color.rgb = CARD
                txt = rows[i - 1][j]
                # numbers and symbols ("Rs 1,250", "-", "9:15") centre in their column; words read from the left
                is_num = plain[i - 1][j] and (any(ch.isdigit() for ch in txt) or not any(ch.isalpha() for ch in txt))
                _cell_text(cell, txt, pt, INK, align=CENTER if is_num else LEFT)
            else:
                cell.fill.fore_color.rgb = CARD
                _cell_text(cell, "", pt, INK)
            _set_borders(cell, LINE2)
    if lay["note"]:
        f.text(x, y + total + 0.1, w, lay["note_h"], lay["note"], lay["npt"], INK2, italic=True)


def split_table(v, lang, w, h):
    """A table too long for one slide, cut into parts that each fit (header repeated; blank rows and note last)."""
    rows = list(v.get("rows") or [])
    for parts in range(2, 5):
        size = -(-len(rows) // parts)
        chunks = [rows[i:i + size] for i in range(0, len(rows), size)] or [[]]
        out = []
        for k, ch in enumerate(chunks):
            part = {kk: vv for kk, vv in v.items() if kk not in ("rows", "blank_rows", "note")}
            part["rows"] = ch
            if k == len(chunks) - 1:
                if v.get("blank_rows"):
                    part["blank_rows"] = v["blank_rows"]
                if v.get("note"):
                    part["note"] = v["note"]
            out.append(part)
        if all(table_plan(p, lang, w, h, pt_min=10, keep_blank=True) for p in out):
            return out
    return None


def need_height(v, lang, w):
    """Height a full-width band needs to show v at a readable size (about 11pt), card title and caption included."""
    t = lambda x: L(x, lang)
    iw = w - 2 * PAD
    extra = 0.34
    if v.get("title"):
        extra += fit.height(t(v["title"]), 13.5, iw, bold=True) + 0.12
    if v.get("caption"):
        extra += min(0.8, fit.height(t(v["caption"]), 10, iw)) + 0.1
    k = v.get("kind")
    if k == "steps":
        items = v.get("items") or []
        inner = iw / max(1, len(items)) - 0.16
        body = 0.55 + max(fit.height(t(s.get("label")), 11.5, inner, bold=True)
                          + (fit.height(t(s.get("sub")), 10, inner) + 0.04 if s.get("sub") else 0) for s in items)
    elif k == "flow":
        nodes = v.get("nodes") or []
        ag = 0.78 if v.get("arrows") else 0.42
        inner = (iw - ag * (len(nodes) - 1)) / max(1, len(nodes)) - 0.24
        body = max(0.9, 0.3 + max(fit.height(t(nd.get("label")), 11.5, inner, bold=True)
                                  + (fit.height(t(nd.get("sub")), 10, inner) if nd.get("sub") else 0) for nd in nodes))
    elif k == "compare":
        cols = v.get("cols") or []
        n = max(1, len(cols))
        inner = (iw - 0.18 * (n - 1)) / n - 0.32
        head = max(fit.height(t(c.get("head")), 12.5, inner, bold=True) for c in cols) + 0.24
        pts = max(fit.paras_height([t(p) for p in c.get("points") or []], 11.5, inner - 0.22, 1.0, 5.75) for c in cols)
        body = head + pts + 0.55
    elif k == "table":
        head, rows, _ = _table_cells(v, lang)
        lay = _table_layout(head, rows, 11, min(3, int(v.get("blank_rows") or 0)), iw)
        body = (lay["total"] if lay else 9.0) + (0.4 if v.get("note") else 0)
    else:
        body = 2.8
    return body + extra


# ------------------------------------------------------------------ mindmap
def v_mindmap(f, v, x, y, w, h):
    branches = list(v.get("branches") or [])
    n = len(branches)
    half = math.ceil(n / 2)
    sides = [branches[:half], branches[half:]]
    center = f.t(v.get("center"))
    cw_c = 2.7
    colw = (w - cw_c - 2 * 0.55) / 2
    inner = colw - 0.32
    gap = 0.18
    chosen = None
    for pt in [x / 2 for x in range(32, 19, -1)]:  # 16 .. 10
        lpt = pt + 1.5
        ok = True
        hs_all = []
        for side in sides:
            hs = []
            for b in side:
                lab = f.t(b.get("label"))
                pts = [f.t(p) for p in (b.get("points") or [])]
                hs.append(fit.height(lab, lpt, inner, bold=True) + 0.08 +
                          fit.paras_height(pts, pt, inner - 0.2, 1.0, 0.3 * pt) + 0.28)
            if sum(hs) + gap * max(0, len(hs) - 1) > h:
                ok = False
            hs_all.append(hs)
        if ok:
            chosen = (pt, lpt, hs_all)
            break
    if not chosen:
        pt, lpt = 10, 11.5
        hs_all = [[h / max(1, len(s)) - gap for _ in s] for s in sides]
        f.overflow("mindmap", center, w, h)
    else:
        pt, lpt, hs_all = chosen
    ccx, ccy = x + w / 2, y + h / 2
    cpt = fit.fit(center, cw_c - 0.3, 1.1, 18, 12, bold=True) or 12
    ch_c = max(1.0, fit.height(center, cpt, cw_c - 0.3, bold=True) + 0.4)
    cards = []
    for si, side in enumerate(sides):
        hs = hs_all[si]
        tot = sum(hs) + gap * max(0, len(hs) - 1)
        cy = y + (h - tot) / 2
        bx = x if si == 0 else x + w - colw
        for bi, b in enumerate(side):
            cards.append((si, b, bx, cy, hs[bi]))
            cy += hs[bi] + gap
    # connectors first, so the cards sit on top of them
    for si, b, bx, by, bh in cards:
        t = tone(b.get("tone"))
        ex = bx + colw if si == 0 else bx
        sx = ccx - cw_c / 2 if si == 0 else ccx + cw_c / 2
        c = f.s.shapes.add_connector(MSO_CONNECTOR.CURVE, E(sx), E(ccy), E(ex), E(by + bh / 2))
        c.line.color.rgb = TONE_FILL[t]
        c.line.width = Pt(1.75)
    cb = f.box(ccx - cw_c / 2, ccy - ch_c / 2, cw_c, ch_c, fill=NAVY, line=None, radius=ch_c / 2)
    f.cv.label(cb, center, cpt, WHITE, bold=True, margin=0.15)
    for si, b, bx, by, bh in cards:
        t = tone(b.get("tone"))
        f.box(bx, by, colw, bh, fill=CARD, line=TONE_FILL[t], lw=1.5, radius=0.12)
        lab = f.t(b.get("label"))
        pts = [f.t(p) for p in (b.get("points") or [])]
        lh = fit.height(lab, lpt, inner, bold=True)
        f.text(bx + 0.16, by + 0.12, inner, lh + 0.02, lab, lpt, TONE_TEXT[t], bold=True)
        f.text(bx + 0.16, by + 0.12 + lh + 0.08, inner, bh - lh - 0.2, pts, pt, INK, gap=0.3 * pt, bullet="•",
               bullet_color=TONE_FILL[t], indent=0.2)


RENDERERS = {"calc": v_calc, "bars": v_bars, "compare": v_compare, "steps": v_steps, "flow": v_flow,
             "line": v_line, "candles": v_candles, "table": v_table, "mindmap": v_mindmap}
